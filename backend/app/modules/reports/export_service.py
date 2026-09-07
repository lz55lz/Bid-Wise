"""将冻结 Markdown 报告导出为不同文件格式，不重新调用模型。"""

from dataclasses import dataclass
from html import escape
from io import BytesIO
from uuid import UUID

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.core.errors import DomainError
from app.modules.identity.service import AuthenticatedUser
from app.modules.reports.service import ProjectReportService


@dataclass(frozen=True, slots=True)
class ReportExport:
    """一个已授权、可直接返回 HTTP 的内存导出文件。"""

    content: bytes
    mime_type: str
    extension: str


class ReportExportService:
    """导出只读取 `READY` 报告的 Markdown 快照，所有格式保持相同事实内容。"""

    def __init__(self, session: AsyncSession, settings: Settings) -> None:
        self._reports = ProjectReportService(session, settings)

    async def export(
        self,
        project_id: UUID,
        report_id: UUID,
        actor: AuthenticatedUser,
        report_format: str,
    ) -> ReportExport:
        """先通过报告服务校验项目授权和报告状态，再渲染请求格式。"""
        report = await self._reports.get(project_id, report_id, actor)
        if report.status != "READY" or not report.content_markdown:
            raise DomainError("REPORT_NOT_READY", "报告尚未生成完成，不能导出", 409)
        if report_format == "md":
            return ReportExport(report.content_markdown.encode("utf-8"), "text/markdown", "md")
        if report_format == "docx":
            return ReportExport(
                self._render_docx(report.content_markdown),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "docx",
            )
        if report_format == "pdf":
            return ReportExport(self._render_pdf(report.content_markdown), "application/pdf", "pdf")
        raise DomainError("VALIDATION_ERROR", "不支持的报告导出格式", 422)

    @staticmethod
    def _render_docx(markdown: str) -> bytes:
        """使用最小 Markdown 子集生成 DOCX，避免将模型输出当作模板执行。"""
        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.size = Pt(10.5)
        for line in markdown.splitlines():
            text = line.strip()
            if not text:
                document.add_paragraph("")
            elif text.startswith("### "):
                document.add_heading(text[4:], level=3)
            elif text.startswith("## "):
                document.add_heading(text[3:], level=2)
            elif text.startswith("# "):
                document.add_heading(text[2:], level=1)
            elif text.startswith(("- ", "* ")):
                document.add_paragraph(text[2:], style="List Bullet")
            else:
                document.add_paragraph(text)
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _render_pdf(markdown: str) -> bytes:
        """使用 ReportLab 内置中文 CID 字体渲染 PDF，避免依赖宿主机字体路径。"""
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        styles = getSampleStyleSheet()
        body = ParagraphStyle(
            "BidWiseBody",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=10,
            leading=16,
            spaceAfter=4,
        )
        headings = {
            "# ": ParagraphStyle("H1Cn", parent=body, fontSize=18, leading=26, spaceBefore=10),
            "## ": ParagraphStyle("H2Cn", parent=body, fontSize=14, leading=22, spaceBefore=8),
            "### ": ParagraphStyle("H3Cn", parent=body, fontSize=12, leading=18, spaceBefore=6),
        }
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
        )
        flowables: list[object] = []
        for line in markdown.splitlines():
            text = line.strip()
            if not text:
                flowables.append(Spacer(1, 4))
                continue
            style = body
            for prefix, heading_style in headings.items():
                if text.startswith(prefix):
                    text = text[len(prefix) :]
                    style = heading_style
                    break
            if text.startswith(("- ", "* ")):
                text = f"• {text[2:]}"
            flowables.append(Paragraph(escape(text), style))
        document.build(flowables)
        return output.getvalue()
